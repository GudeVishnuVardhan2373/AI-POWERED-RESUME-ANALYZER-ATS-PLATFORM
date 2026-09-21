import re
import os
from typing import Dict, Any, List, Optional
from pypdf import PdfReader
from docx import Document

EMAIL_REGEX = r'[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+'
PHONE_REGEX = r'(?:\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}|\+?\d{10,13}'
LINKEDIN_REGEX = r'(?:https?:\/\/)?(?:www\.)?linkedin\.com\/in\/[a-zA-Z0-9_-]+'
GITHUB_REGEX = r'(?:https?:\/\/)?(?:www\.)?github\.com\/[a-zA-Z0-9_-]+'
URL_REGEX = r'https?:\/\/(?:www\.)?[-a-zA-Z0-9@:%._\+~#=]{1,256}\.[a-zA-Z0-9()]{1,6}\b(?:[-a-zA-Z0-9()@:%_\+.~#?&//=]*)'

DEGREE_PATTERNS = [
    r'\b(?:B\.?Tech|B\.?E|B\.?S|B\.?Sc|BCA|Bachelor of [A-Za-z\s]+)\b',
    r'\b(?:M\.?Tech|M\.?E|M\.?S|M\.?Sc|MCA|MBA|Master of [A-Za-z\s]+)\b',
    r'\b(?:Ph\.?D|Doctorate)\b',
    r'\b(?:Diploma in [A-Za-z\s]+)\b'
]

EXPERIENCE_PATTERNS = [
    r'(\d+(?:\.\d+)?)\s*\+?\s*(?:years?|yrs?)(?:\s+of)?\s+(?:experience|exp)',
    r'(?:experience|exp)\s*(?::|-)?\s*(\d+(?:\.\d+)?)\s*\+?\s*(?:years?|yrs?)'
]

def extract_text_from_file(file_path: str) -> str:
    """Extract full raw text from PDF, DOCX, or plain text file."""
    ext = os.path.splitext(file_path)[1].lower()
    text = ""
    
    if ext == ".pdf":
        try:
            reader = PdfReader(file_path)
            pages = []
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    pages.append(extracted)
            text = "\n".join(pages)
        except Exception as e:
            text = f"Error reading PDF: {str(e)}"
            
    elif ext in [".docx", ".doc"]:
        try:
            doc = Document(file_path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            paragraphs.append(cell.text.strip())
            text = "\n".join(paragraphs)
        except Exception as e:
            text = f"Error reading DOCX: {str(e)}"
            
    elif ext in [".txt", ".rtf", ".md"]:
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()
        except Exception as e:
            text = f"Error reading text file: {str(e)}"
    else:
        # Fallback raw read
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()
        except Exception as e:
            text = f"Unsupported file type: {str(e)}"

    return text.strip()


def parse_candidate_profile(raw_text: str, filename: str) -> Dict[str, Any]:
    """Extract contact information, candidate name, links, education, and experience from resume text."""
    # 1. Email extraction
    emails = re.findall(EMAIL_REGEX, raw_text)
    email = emails[0] if emails else None

    # 2. Phone extraction
    phones = re.findall(PHONE_REGEX, raw_text)
    # Filter out short or false positives
    valid_phones = [p.strip() for p in phones if len(re.sub(r'\D', '', p)) >= 10]
    phone = valid_phones[0] if valid_phones else None

    # 3. Links extraction
    links = []
    linkedin_matches = re.findall(LINKEDIN_REGEX, raw_text, re.IGNORECASE)
    github_matches = re.findall(GITHUB_REGEX, raw_text, re.IGNORECASE)
    links.extend(linkedin_matches)
    links.extend(github_matches)
    # Deduplicate
    links = list(dict.fromkeys(links))

    # 4. Education extraction
    education = []
    for pat in DEGREE_PATTERNS:
        matches = re.findall(pat, raw_text, re.IGNORECASE)
        for m in matches:
            if m.strip() and m.strip() not in education:
                education.append(m.strip())

    # 5. Experience extraction
    years_exp = 0.0
    for pat in EXPERIENCE_PATTERNS:
        matches = re.findall(pat, raw_text, re.IGNORECASE)
        if matches:
            try:
                found_years = float(matches[0])
                if found_years > years_exp and found_years < 45:
                    years_exp = found_years
            except ValueError:
                pass

    # If not explicitly stated as "X years", detect date intervals like 2021 - 2024 or 2021 - Present
    if years_exp == 0.0:
        year_spans = re.findall(r'\b(20\d{2}|19\d{2})\s*(?:-|to|–)\s*(20\d{2}|Present|Current)\b', raw_text, re.IGNORECASE)
        total_span_years = 0
        current_year = 2026
        for start_str, end_str in year_spans:
            try:
                start = int(start_str)
                end = current_year if end_str.lower() in ["present", "current"] else int(end_str)
                span = max(0, end - start)
                if span < 30:
                    total_span_years += span
            except Exception:
                pass
        if total_span_years > 0:
            years_exp = min(float(total_span_years), 30.0)

    # 6. Candidate Name heuristics
    # Look at the first 5 non-empty lines
    lines = [l.strip() for l in raw_text.splitlines() if l.strip()]
    candidate_name = None
    blacklist_words = ["resume", "curriculum", "vitae", "cv", "profile", "contact", "summary", "experience", "education"]

    for line in lines[:5]:
        # Filter out email or phone lines
        if "@" in line or any(d in line for d in "0123456789+"):
            continue
        cleaned = re.sub(r'[^a-zA-Z\s]', '', line).strip()
        words = cleaned.split()
        if 2 <= len(words) <= 4:
            if not any(bw in cleaned.lower() for bw in blacklist_words):
                candidate_name = cleaned.title()
                break

    if not candidate_name:
        # Fallback to filename without extension
        base_name = os.path.splitext(filename)[0]
        cleaned_base = re.sub(r'[_\-\d]', ' ', base_name).strip()
        candidate_name = cleaned_base.title() if cleaned_base else "Candidate"

    return {
        "candidate_name": candidate_name,
        "email": email,
        "phone": phone,
        "links": links,
        "education": education,
        "experience_years": years_exp
    }
